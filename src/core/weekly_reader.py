"""Leitura da planilha modelo do relatório semanal (ex: 'Modelo Hotel - ... - Venda Semanal.xlsx').

Estrutura esperada (por posição ou nome da aba):
  1. Capa
  2. Emissões por Canal      (tabela dinâmica, cabeçalho na linha 3, hotel em B1)
  3. Emissões por Categoria  (tabela dinâmica, cabeçalho na linha 3)
  4. Evolução Semanal        (tabela dinâmica, cabeçalho na linha 3)
  5. Base de Dados           (reservas emitidas na semana, cabeçalho na linha 1)
  6. Histórico de Venda Semanal (uma linha por semana, cabeçalho na linha 1)
"""
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date
from typing import List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook


PT_MONTHS = {
    "janeiro": 1, "jan": 1,
    "fevereiro": 2, "fev": 2,
    "março": 3, "marco": 3, "mar": 3,
    "abril": 4, "abr": 4,
    "maio": 5, "mai": 5,
    "junho": 6, "jun": 6,
    "julho": 7, "jul": 7,
    "agosto": 8, "ago": 8,
    "setembro": 9, "set": 9,
    "outubro": 10, "out": 10,
    "novembro": 11, "nov": 11,
    "dezembro": 12, "dez": 12,
}


def _parse_week_end_date(text: str) -> Optional[Tuple[int, int, Optional[int]]]:
    s = str(text).strip().lower()
    m = re.search(r'(\d{1,2})\s+a\s+(\d{1,2})\s+de\s+(?:de\s+)?([a-zçã]+)(?:\s+de\s+(\d{4}))?', s)
    if not m:
        return None
    end_day = int(m.group(2))
    month_name = m.group(3)
    explicit_year = int(m.group(4)) if m.group(4) else None
    month = PT_MONTHS.get(month_name)
    if not month:
        return None
    return (end_day, month, explicit_year)


def filter_evolution_last_quarter(df: pd.DataFrame, base_year: Optional[int] = None) -> pd.DataFrame:
    """Filtra a tabela de evolução semanal para conter apenas o último trimestre (D-3 meses)."""
    if df is None or df.empty or len(df) <= 13:
        return df

    year = base_year or 2026
    first_col = df.columns[0]

    # Detecta se está ordenado de forma decrescente (mais recente primeiro)
    seq_col = next((c for c in df.columns if "sequ" in _norm(c)), None)
    is_desc = True
    if seq_col is not None:
        seq_vals = pd.to_numeric(df[seq_col], errors="coerce").dropna()
        if len(seq_vals) >= 2:
            is_desc = bool(seq_vals.iloc[0] > seq_vals.iloc[-1])

    parsed_info = [_parse_week_end_date(v) for v in df[first_col]]
    valid_count = sum(1 for p in parsed_info if p is not None)
    if valid_count < len(df) * 0.5:
        return df.iloc[:14].copy() if is_desc else df.iloc[-14:].copy()

    dates = []
    curr_year = year
    prev_month = None
    for p in parsed_info:
        if p is None:
            dates.append(None)
            continue
        end_day, month, explicit_year = p
        if explicit_year:
            curr_year = explicit_year
        elif prev_month is not None:
            if is_desc and month > prev_month:
                curr_year -= 1
            elif not is_desc and month < prev_month:
                curr_year += 1
        prev_month = month
        try:
            d = date(curr_year, month, end_day)
        except ValueError:
            d = None
        dates.append(d)

    valid_dates = [d for d in dates if d is not None]
    if not valid_dates:
        return df.iloc[:14].copy() if is_desc else df.iloc[-14:].copy()

    max_date = max(valid_dates)
    cutoff_date = (pd.Timestamp(max_date) - pd.DateOffset(months=3)).date()

    mask = [d is not None and d >= cutoff_date for d in dates]
    filtered = df[mask].copy()
    if filtered.empty:
        return df.iloc[:14].copy() if is_desc else df.iloc[-14:].copy()
    return filtered


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return text.lower().strip()


def _is_vazio(val) -> bool:
    """Verifica se o valor representa uma linha/célula vazia ou rótulo '(vazio)' de tabela dinâmica."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return True
    s = _norm(val)
    clean = re.sub(r'[\(\)\s\[\]<>]', '', s)
    if not clean or clean in ('none', 'nan', 'null', '-'):
        return True
    if 'vazi' in clean or 'blank' in clean:
        return True
    return False


def _is_emissao_por_uh(col_name: str) -> bool:
    """Detecta variações do cabeçalho de emissão por UH (ex: 'Média de Emissão por UH', 'Emissão por UH')."""
    n = _norm(col_name)
    return "emissao por uh" in n or "por uh" in n or bool(re.search(r'emiss.*uh', n))



@dataclass
class WeeklyData:
    """Conteúdo extraído da planilha modelo semanal."""
    hotel: str = ""
    year: Optional[int] = None
    week_label: str = ""
    by_channel: pd.DataFrame = field(default_factory=pd.DataFrame)
    by_category: pd.DataFrame = field(default_factory=pd.DataFrame)
    evolution: pd.DataFrame = field(default_factory=pd.DataFrame)
    base: pd.DataFrame = field(default_factory=pd.DataFrame)
    history: pd.DataFrame = field(default_factory=pd.DataFrame)

    @property
    def reference(self) -> str:
        if not self.week_label:
            return ""
        return f"Semana de {self.week_label}" + (f" de {self.year}" if self.year else "")


class WeeklyWorkbookReader:
    """Lê as abas relevantes da planilha semanal (modo streaming, pois os arquivos são pesados)."""

    # (chave, trecho do nome normalizado, posição fallback 0-based)
    _SHEETS = {
        "channel": ("canal", 1),
        "category": ("categoria", 2),
        "evolution": ("evolu", 3),
        "base": ("base", 4),
        "history": ("histor", 5),
    }

    def read(self, file_path: str) -> WeeklyData:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            names = wb.sheetnames
            sheets = {key: self._resolve_sheet(names, key) for key in self._SHEETS}

            hotel = self._read_hotel(wb[sheets["channel"]]) if sheets["channel"] else ""
            data = WeeklyData(hotel=hotel)
            if sheets["channel"]:
                data.by_channel = self._read_block(wb[sheets["channel"]], header_row=3)
            if sheets["category"]:
                data.by_category = self._read_block(wb[sheets["category"]], header_row=3)
            if sheets["evolution"]:
                df_ev = self._read_block(wb[sheets["evolution"]], header_row=3)
                if not df_ev.empty:
                    first_col = df_ev.columns[0]
                    df_ev = df_ev[~df_ev[first_col].apply(_is_vazio)].reset_index(drop=True)
                    uh_cols = [c for c in df_ev.columns if _is_emissao_por_uh(c)]
                    if uh_cols:
                        df_ev = df_ev.drop(columns=uh_cols)
                data.evolution = df_ev
            if sheets["base"]:
                data.base = self._read_base(wb[sheets["base"]])
            if sheets["history"]:
                data.history = self._read_block(wb[sheets["history"]], header_row=1)
        finally:
            wb.close()

        if data.by_channel.empty and data.by_category.empty and data.evolution.empty:
            raise ValueError("Nenhuma tabela encontrada nas abas 2 a 4 do arquivo. Verifique se é a planilha modelo semanal.")

        self._fill_context(data)
        return data

    # ------------------------------------------------------------------ helpers

    def _resolve_sheet(self, names: List[str], key: str) -> Optional[str]:
        token, fallback_idx = self._SHEETS[key]
        for n in names:
            if token in _norm(n):
                return n
        return names[fallback_idx] if fallback_idx < len(names) else None

    @staticmethod
    def _read_hotel(ws) -> str:
        for row in ws.iter_rows(min_row=1, max_row=2, values_only=True):
            cells = [c for c in row if c is not None]
            if len(cells) >= 2 and "hotel" in _norm(cells[0]):
                return str(cells[1]).strip()
        return ""

    @staticmethod
    def _read_block(ws, header_row: int) -> pd.DataFrame:
        """Lê um bloco contíguo (cabeçalho + linhas até a primeira linha vazia)."""
        header, rows = None, []
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx < header_row:
                continue
            if idx == header_row:
                header = [str(c).strip() if c is not None else "" for c in row]
                last = max((i for i, h in enumerate(header) if h), default=-1)
                header = header[:last + 1]
                continue
            if row is None or row[0] is None:
                break
            rows.append(list(row[:len(header)]))
        if not header:
            return pd.DataFrame()
        return pd.DataFrame(rows, columns=header)

    @staticmethod
    def _read_base(ws) -> pd.DataFrame:
        df = WeeklyWorkbookReader._read_block(ws, header_row=1)
        if df.empty:
            return df
        # Remove eventuais linhas sem hotel/canal (a tabela do Excel vai até a última linha da planilha)
        if "Canal" in df.columns:
            df = df.dropna(subset=["Canal"])
        for col in ("Data de Compra", "Check-in", "Check-out"):
            if col in df.columns:
                df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=True)
        for col in ("Valor da Reserva", "LOS", "ADS", "ADR", "Adultos", "Crianças"):
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors="coerce")
        if "ADS" in df.columns and "Check-in" in df.columns and "Data de Compra" in df.columns:
            calc = (df["Check-in"] - df["Data de Compra"]).dt.days
            df["ADS"] = df["ADS"].fillna(calc)
        if "ADR" in df.columns and "Valor da Reserva" in df.columns and "LOS" in df.columns:
            calc = df["Valor da Reserva"] / df["LOS"].replace(0, pd.NA)
            df["ADR"] = df["ADR"].fillna(calc)
        return df.reset_index(drop=True)

    @staticmethod
    def _fill_context(data: WeeklyData) -> None:
        """Descobre o ano e o rótulo da semana mais recente."""
        if not data.base.empty and "Data de Compra" in data.base.columns:
            max_dt = data.base["Data de Compra"].max()
            if pd.notna(max_dt):
                data.year = int(max_dt.year)

        ev = data.evolution
        if ev.empty:
            return
        label_col = ev.columns[0]
        valid_ev = ev[~ev[label_col].apply(_is_vazio)]
        if valid_ev.empty:
            return
        seq_col = next((c for c in valid_ev.columns if "sequ" in _norm(c)), None)
        if seq_col is not None:
            seq = pd.to_numeric(valid_ev[seq_col], errors="coerce")
            if seq.notna().any():
                data.week_label = str(valid_ev.loc[seq.idxmax(), label_col]).strip()
                return
        data.week_label = str(valid_ev.iloc[0][label_col]).strip()
